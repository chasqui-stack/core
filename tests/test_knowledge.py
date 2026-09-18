"""Sprint 15.1 acceptance: knowledge module — extract, chunk, embed, admin routes.

Embeddings are faked with deterministic vectors so pgvector computes REAL
cosine distances (the test DB has the extension). Fixture files are built
in memory — no binary blobs in the repo, no network.
"""

import io
import uuid
from datetime import timedelta

import pytest
from sqlmodel import select

from app.core.config import settings
from app.modules.knowledge import service as knowledge_service
from app.modules.knowledge.extract import (
    EmptyText,
    UnsupportedType,
    extract_text,
    validate_type,
)
from app.modules.knowledge.models import Document, DocumentChunk
from app.services.admin_service import create_admin_access_token

BASE = "/admin/modules/knowledge"


def vec(*head: float) -> list[float]:
    """A settings.embedding_dim-wide vector with `head` as leading components."""
    v = [0.0] * settings.embedding_dim
    v[: len(head)] = head
    return v


class FakeEmbeddings:
    """Texts mentioning 'refund' point one way, everything else another."""

    def __init__(self) -> None:
        self.document_calls: list[list[str]] = []
        self.fail = False

    def _vector(self, text: str) -> list[float]:
        return vec(1.0, 0.0) if "refund" in text.lower() else vec(0.0, 1.0)

    async def aembed_query(self, text: str) -> list[float]:
        if self.fail:
            raise RuntimeError("embeddings provider down")
        return self._vector(text)

    async def aembed_documents(self, texts: list[str]) -> list[list[float]]:
        if self.fail:
            raise RuntimeError("embeddings provider down")
        self.document_calls.append(list(texts))
        return [self._vector(t) for t in texts]


@pytest.fixture
def fake_embeddings(monkeypatch) -> FakeEmbeddings:
    import app.core.embeddings as embeddings_mod

    fake = FakeEmbeddings()
    monkeypatch.setattr(embeddings_mod, "get_embeddings", lambda: fake)
    return fake


@pytest.fixture(autouse=True)
def _job_sessions(monkeypatch, session_factory):
    """Background jobs open their own sessions — bind them to the test txn."""
    import app.db.session as db_session

    monkeypatch.setattr(db_session, "async_session_factory", session_factory)


@pytest.fixture
def admin_headers() -> dict:
    token = create_admin_access_token(uuid.uuid4(), "admin@test.local", "super_admin")
    return {"Authorization": f"Bearer {token}"}


# --- in-memory fixture files -------------------------------------------------


def make_pdf(text: str | None) -> bytes:
    """A one-page PDF; `text=None` mimics a scan (a page with no text layer)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref)
    )
    return out.getvalue()


def make_docx(paragraph: str, table_row: list[str]) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph(paragraph)
    table = doc.add_table(rows=1, cols=len(table_row))
    for cell, value in zip(table.rows[0].cells, table_row):
        cell.text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


LONG_TEXT = "\n\n".join(
    f"Section {i}. " + " ".join(f"word{i}x{j}" for j in range(60)) for i in range(12)
)


# --- extraction --------------------------------------------------------------


def test_extracts_every_supported_type():
    assert "Refunds take 5 days" in extract_text("policy.pdf", make_pdf("Refunds take 5 days"))

    docx = extract_text("prices.docx", make_docx("Price list", ["Basic", "10 USD"]))
    assert "Price list" in docx and "Basic | 10 USD" in docx

    html = extract_text(
        "page.html",
        b"<html><head><style>p{color:red}</style><script>var x=1</script></head>"
        b"<body><h1>Shipping</h1><p>Two days.</p></body></html>",
    )
    assert "Shipping" in html and "Two days." in html
    assert "color" not in html and "var x" not in html

    assert extract_text("notes.txt", "hola ñandú".encode()) == "hola ñandú"
    assert extract_text("README.md", b"\xef\xbb\xbf# Title\nBody") == "# Title\nBody"


def test_docx_keeps_tables_in_document_order():
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Prices")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Basic", "10 USD"
    doc.add_paragraph("Returns policy")
    out = io.BytesIO()
    doc.save(out)

    # The table stays under its heading — not appended after the last paragraph
    assert extract_text("prices.docx", out.getvalue()) == (
        "Prices\nBasic | 10 USD\nReturns policy"
    )


def test_scanned_pdf_and_blank_files_have_no_text():
    with pytest.raises(EmptyText, match="no extractable text"):
        extract_text("scan.pdf", make_pdf(None))
    with pytest.raises(EmptyText):
        extract_text("blank.txt", b"  \n\x00 ")


def test_type_whitelist():
    validate_type("manual.PDF", "application/pdf")
    validate_type("notes.md", "application/octet-stream")  # what curl sends
    validate_type("notes.md", None)
    with pytest.raises(UnsupportedType):
        validate_type("legacy.doc", "application/msword")
    with pytest.raises(UnsupportedType):
        validate_type("no-extension", "text/plain")
    with pytest.raises(UnsupportedType):
        validate_type("fake.pdf", "image/png")


def test_chunks_overlap():
    chunks = knowledge_service.split_text("manual.txt", LONG_TEXT)
    assert len(chunks) > 2
    assert all(len(c) <= knowledge_service.CHUNK_SIZE for c in chunks)
    # consecutive chunks share text: the head of each one appears in the previous
    shared = sum(
        1 for prev, nxt in zip(chunks, chunks[1:]) if nxt.split()[0] in prev
    )
    assert shared >= 1


# --- service -----------------------------------------------------------------


async def upload(session, filename: str, data: bytes) -> Document:
    document = await knowledge_service.create_document(
        session, filename=filename, mime_type=None, data=data
    )
    await session.commit()
    return document


async def test_process_happy_path(session, session_factory, fake_embeddings):
    document = await upload(session, "manual.txt", LONG_TEXT.encode())
    assert document.status == "pending"

    await knowledge_service.process_document(
        document.id, LONG_TEXT.encode(), session_factory
    )

    await session.refresh(document)
    assert document.status == "ready"
    assert document.error_detail is None
    assert document.content_text == LONG_TEXT

    chunks = (
        await session.exec(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document.id)
            .order_by(DocumentChunk.seq)
        )
    ).all()
    assert document.chunk_count == len(chunks) > 2
    assert [c.seq for c in chunks] == list(range(len(chunks)))
    assert all(c.embedding is not None for c in chunks)
    assert len(fake_embeddings.document_calls) == 1  # ONE batched embed call


async def test_scanned_pdf_ends_in_error(session, session_factory, fake_embeddings):
    data = make_pdf(None)
    document = await upload(session, "scan.pdf", data)

    await knowledge_service.process_document(document.id, data, session_factory)

    await session.refresh(document)
    assert document.status == "error"
    assert document.error_detail == "no extractable text"
    assert document.chunk_count == 0
    assert fake_embeddings.document_calls == []


async def test_embeddings_outage_then_reprocess_recovers(
    session, session_factory, fake_embeddings
):
    document = await upload(session, "manual.txt", LONG_TEXT.encode())

    fake_embeddings.fail = True
    await knowledge_service.process_document(
        document.id, LONG_TEXT.encode(), session_factory
    )
    await session.refresh(document)
    assert document.status == "error"
    assert "embeddings provider down" in document.error_detail
    assert document.content_text == LONG_TEXT  # kept — reprocess needs it
    leftovers = await session.exec(
        select(DocumentChunk).where(DocumentChunk.document_id == document.id)
    )
    assert leftovers.all() == []  # never half-indexed

    fake_embeddings.fail = False
    await knowledge_service.process_document(document.id, None, session_factory)
    await session.refresh(document)
    assert document.status == "ready"
    assert document.error_detail is None
    assert document.chunk_count > 0


async def test_reprocess_replaces_chunks(session, session_factory, fake_embeddings):
    document = await upload(session, "manual.txt", LONG_TEXT.encode())
    await knowledge_service.process_document(
        document.id, LONG_TEXT.encode(), session_factory
    )
    await session.refresh(document)
    first_count = document.chunk_count

    await knowledge_service.process_document(document.id, None, session_factory)

    total = await session.exec(
        select(DocumentChunk).where(DocumentChunk.document_id == document.id)
    )
    assert len(total.all()) == first_count  # replaced, not appended


async def test_duplicate_upload_is_rejected(session):
    await upload(session, "a.txt", b"same bytes")
    with pytest.raises(knowledge_service.DuplicateDocument):
        await knowledge_service.create_document(
            session, filename="renamed.txt", mime_type=None, data=b"same bytes"
        )


async def test_search_ranks_and_skips_unready(session, session_factory, fake_embeddings):
    refunds = await upload(session, "refunds.txt", b"Refund policy: 30 days with receipt.")
    hours = await upload(session, "hours.txt", b"Opening hours: Mon-Fri 9-18.")
    for doc, data in ((refunds, b"Refund policy: 30 days with receipt."),
                      (hours, b"Opening hours: Mon-Fri 9-18.")):
        await knowledge_service.process_document(doc.id, data, session_factory)

    hits = await knowledge_service.search(session, "refund", top_k=5, min_similarity=0.0)
    assert [doc.filename for _, doc, _ in hits] == ["refunds.txt", "hours.txt"]
    assert hits[0][2] == pytest.approx(1.0)

    floored = await knowledge_service.search(session, "refund", min_similarity=0.5)
    assert [doc.filename for _, doc, _ in floored] == ["refunds.txt"]

    # a document that is not `ready` is invisible, even with chunks in place
    await session.refresh(refunds)
    refunds.status = "error"
    session.add(refunds)
    await session.commit()
    hidden = await knowledge_service.search(session, "refund", min_similarity=0.5)
    assert hidden == []


async def test_search_degrades_on_embeddings_outage(session, fake_embeddings):
    fake_embeddings.fail = True
    assert await knowledge_service.search(session, "anything") == []


def test_stale_jobs_stop_being_busy():
    document = Document(
        filename="a.txt", mime_type="text/plain", size_bytes=1, content_sha256="x"
    )
    document.status = "processing"
    assert knowledge_service.is_busy(document)
    document.updated_at -= knowledge_service.STALE_AFTER + timedelta(seconds=1)
    assert not knowledge_service.is_busy(document)  # crashed job — reprocess may take over
    document.status = "ready"
    assert not knowledge_service.is_busy(document)


# --- admin routes ------------------------------------------------------------


async def test_routes_require_auth(client):
    response = await client.get(f"{BASE}/documents")
    assert response.status_code == 401


async def test_upload_flow(client, session, fake_embeddings, admin_headers):
    pdf = make_pdf("Refunds take 5 days")
    created = await client.post(
        f"{BASE}/documents",
        files={"file": ("policy.pdf", pdf, "application/pdf")},
        headers=admin_headers,
    )
    assert created.status_code == 202
    body = created.json()
    assert body["status"] == "pending" and body["filename"] == "policy.pdf"
    assert "content_text" not in body

    # the background job ran after the response
    listed = await client.get(f"{BASE}/documents", headers=admin_headers)
    (doc,) = listed.json()
    assert doc["status"] == "ready" and doc["chunk_count"] == 1
    assert doc["size_bytes"] == len(pdf)

    duplicate = await client.post(
        f"{BASE}/documents",
        files={"file": ("again.pdf", pdf, "application/pdf")},
        headers=admin_headers,
    )
    assert duplicate.status_code == 409

    hits = await client.get(
        f"{BASE}/search", params={"q": "refund"}, headers=admin_headers
    )
    assert hits.status_code == 200
    (hit,) = hits.json()
    assert hit["filename"] == "policy.pdf" and hit["seq"] == 0
    assert "Refunds take 5 days" in hit["content"]
    assert hit["similarity"] == pytest.approx(1.0)

    deleted = await client.delete(f"{BASE}/documents/{doc['id']}", headers=admin_headers)
    assert deleted.status_code == 204
    chunks = await session.exec(select(DocumentChunk))
    assert chunks.all() == []  # FK cascade
    missing = await client.delete(f"{BASE}/documents/{doc['id']}", headers=admin_headers)
    assert missing.status_code == 404


async def test_upload_rejections(client, monkeypatch, admin_headers):
    legacy = await client.post(
        f"{BASE}/documents",
        files={"file": ("old.doc", b"binary", "application/msword")},
        headers=admin_headers,
    )
    assert legacy.status_code == 415
    assert ".doc" in legacy.json()["detail"]

    monkeypatch.setattr(knowledge_service, "MAX_UPLOAD_BYTES", 16)
    too_big = await client.post(
        f"{BASE}/documents",
        files={"file": ("big.txt", b"x" * 17, "text/plain")},
        headers=admin_headers,
    )
    assert too_big.status_code == 413

    empty = await client.post(
        f"{BASE}/documents",
        files={"file": ("empty.txt", b"", "text/plain")},
        headers=admin_headers,
    )
    assert empty.status_code == 400


async def test_upload_survives_embeddings_outage_and_reprocess_recovers(
    client, fake_embeddings, admin_headers
):
    fake_embeddings.fail = True
    created = await client.post(
        f"{BASE}/documents",
        files={"file": ("manual.txt", LONG_TEXT.encode(), "text/plain")},
        headers=admin_headers,
    )
    assert created.status_code == 202  # an outage never 500s the upload
    doc_id = created.json()["id"]

    (doc,) = (await client.get(f"{BASE}/documents", headers=admin_headers)).json()
    assert doc["status"] == "error"
    assert "embeddings provider down" in doc["error_detail"]
    assert doc["can_reprocess"] is True  # the text was extracted and kept

    fake_embeddings.fail = False
    again = await client.post(
        f"{BASE}/documents/{doc_id}/reprocess", headers=admin_headers
    )
    assert again.status_code == 202
    (doc,) = (await client.get(f"{BASE}/documents", headers=admin_headers)).json()
    assert doc["status"] == "ready" and doc["chunk_count"] > 2


async def test_reprocess_conflicts(client, session, admin_headers):
    busy = Document(
        filename="busy.txt",
        mime_type="text/plain",
        size_bytes=3,
        content_sha256="busy",
        content_text="abc",
        status="processing",
    )
    textless = Document(
        filename="scan.pdf",
        mime_type="application/pdf",
        size_bytes=3,
        content_sha256="scan",
        status="error",
        error_detail="no extractable text",
    )
    session.add(busy)
    session.add(textless)
    await session.commit()

    listed = (await client.get(f"{BASE}/documents", headers=admin_headers)).json()
    flags = {d["filename"]: d["can_reprocess"] for d in listed}
    assert flags == {"busy.txt": True, "scan.pdf": False}

    for document in (busy, textless):
        response = await client.post(
            f"{BASE}/documents/{document.id}/reprocess", headers=admin_headers
        )
        assert response.status_code == 409

    missing = await client.post(
        f"{BASE}/documents/{uuid.uuid4()}/reprocess", headers=admin_headers
    )
    assert missing.status_code == 404
