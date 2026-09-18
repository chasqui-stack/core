"""Text extraction — bytes in, plain text out, one light dependency per type.

Synchronous and CPU-bound on purpose: callers run it off the event loop
(`asyncio.to_thread`). No OCR — a scanned PDF yields no text and raises
EmptyText.
"""

import io
from pathlib import PurePath

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".html"}

# Browsers and CLIs disagree on MIME types (curl sends octet-stream for .md),
# so the extension is the gate; the declared type only has to be plausible.
_GENERIC_MIME_TYPES = {"", "application/octet-stream"}
_MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    },
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/x-markdown", "text/plain"},
    ".html": {"text/html", "application/xhtml+xml"},
}


class UnsupportedType(Exception):
    """The file's extension or declared MIME type is not on the whitelist."""


class EmptyText(Exception):
    """Extraction worked but found no text (e.g. a scanned PDF)."""

    def __init__(self) -> None:
        super().__init__("no extractable text")


def extension_of(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def validate_type(filename: str, mime_type: str | None) -> None:
    ext = extension_of(filename)
    if ext not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise UnsupportedType(
            f"Unsupported file type '{ext or filename}' — allowed: {allowed}"
        )
    declared = (mime_type or "").split(";")[0].strip().lower()
    if declared not in _GENERIC_MIME_TYPES and declared not in _MIME_TYPES[ext]:
        raise UnsupportedType(
            f"Content type '{declared}' does not match a '{ext}' file"
        )


def _decode(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(p for p in pages if p)


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs]
    for table in doc.tables:  # price lists and catalogs live in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(p for p in parts if p)


def _extract_html(data: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    lines = (line.strip() for line in soup.get_text(separator="\n").splitlines())
    return "\n".join(line for line in lines if line)


def extract_text(filename: str, data: bytes) -> str:
    """Plain text of a whitelisted file. Raises UnsupportedType / EmptyText."""
    ext = extension_of(filename)
    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".html":
        text = _extract_html(data)
    elif ext in (".txt", ".md"):
        text = _decode(data)
    else:
        raise UnsupportedType(f"Unsupported file type '{ext or filename}'")

    text = text.replace("\x00", "").strip()  # Postgres text rejects NUL bytes
    if not text:
        raise EmptyText()
    return text
